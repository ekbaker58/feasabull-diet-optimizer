import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

d = json.load(open('foods.json'))
FOODS, NUTS = d['foods'], d['nutrients']
res = json.load(open('js_results.json'))['base']
xopt = dict(zip(res['foods'], res['x']))
NF = len(FOODS)  # 36

BLUE, BLACK, GREEN = '0000FF', '000000', '008000'
def font(color=BLACK, bold=False, size=10, italic=False): return Font(name='Arial', color=color, bold=bold, size=size, italic=italic)
YELLOW = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor='1F3864')
SUB = PatternFill('solid', fgColor='D9E1F2')
GREY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

def put(ws, ref, val, color=BLACK, bold=False, fmt=None, fill=None, align=None, border=False, italic=False, size=10):
    c = ws[ref]; c.value = val; c.font = font(color, bold, size, italic)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = BOX
    return c

def header(ws, ref, text):
    put(ws, ref, text, 'FFFFFF', True, fill=HEAD, align=CENTER, border=True)

wb = Workbook()
# ------------------------------------------------------------------ Food Data
fd = wb.active; fd.title = 'Food Data'
put(fd, 'A1', 'Food Data — prices, edible yields, and nutrients per 100 g edible portion', bold=True, size=14)
put(fd, 'A2', 'Prices: BLS Average Price Data, U.S. city average, September 2026.  Nutrients: USDA National Nutrient Database for Standard Reference, Release 28 (SR28).  See "Notes & Sources".', italic=True)
put(fd, 'A3', 'Legend:  blue = input data (edit freely)   black = formula   yellow = price still needed (BLS no longer publishes it) — enter a local price and the food joins the model.', italic=True)
cols = ['#','Food','Category','BLS item (exact name)','Price ($)','Price unit','Grams per price unit (as bought)','Edible yield','Edible grams per unit',
        'Cost per 100 g edible ($)','Serving (g)','Serving','Cost per serving ($)','Max servings / day']
cols += [f"{n['label']} ({n['unit']})" for n in NUTS]
cols += ['Protein per $ (g)','Calories per $ (kcal)','USDA NDB No.','USDA food description','USDA refuse','Notes']
HR = 5
for i, h in enumerate(cols, 1): header(fd, f'{L(i)}{HR}', h)
fd.row_dimensions[HR].height = 54
NUTCOL = {n['key']: L(15 + i) for i, n in enumerate(NUTS)}   # O..Y
FIRST = HR + 1
for k, f in enumerate(FOODS):
    r = FIRST + k
    put(fd, f'A{r}', k + 1, border=True)
    put(fd, f'B{r}', f['name'], bold=True, border=True)
    put(fd, f'C{r}', f['category'], border=True)
    put(fd, f'D{r}', f['bls_item'], border=True)
    if f['price'] is None:
        put(fd, f'E{r}', None, BLUE, fmt='$0.000', fill=YELLOW, border=True)
    else:
        put(fd, f'E{r}', f['price'], BLUE, fmt='$0.000', border=True)
    put(fd, f'F{r}', f['price_unit'], border=True)
    g = {'eggs': '=12*50/0.88', 'milk': '=16*244', 'ice_cream': '=8*132'}.get(f['id'], f['purchase_g'])
    put(fd, f'G{r}', g, BLUE, fmt='#,##0.0', border=True)
    put(fd, f'H{r}', f['edible_yield'], BLUE, fmt='0%', border=True)
    put(fd, f'I{r}', f'=G{r}*H{r}', fmt='#,##0.0', border=True)
    put(fd, f'J{r}', f'=IF(E{r}="","",E{r}/I{r}*100)', fmt='$0.000', border=True)
    put(fd, f'K{r}', f['serving_g'], BLUE, fmt='0.0', border=True)
    put(fd, f'L{r}', f['serving_desc'], border=True)
    put(fd, f'M{r}', f'=IF(J{r}="","",J{r}*K{r}/100)', fmt='$0.00', border=True)
    put(fd, f'N{r}', f['max_servings'], BLUE, border=True)
    for n in NUTS:
        put(fd, f"{NUTCOL[n['key']]}{r}", f['per100g'][n['key']], BLUE, fmt='0.00', border=True)
    put(fd, f'Z{r}', f'=IF(J{r}="","",P{r}/J{r})', fmt='0.0', border=True)
    put(fd, f'AA{r}', f'=IF(J{r}="","",O{r}/J{r})', fmt='#,##0', border=True)
    put(fd, f'AB{r}', f['ndb'], border=True)
    put(fd, f'AC{r}', f['usda_desc'], border=True)
    put(fd, f'AD{r}', f"{f['refuse_pct']:.0f}% {f['refuse_desc']}".strip() if f['refuse_pct'] else '0%', border=True)
    put(fd, f'AE{r}', f['note'], border=True)
LAST = FIRST + NF - 1
fd['G' + str(FIRST + [f['id'] for f in FOODS].index('eggs'))].comment = Comment('12 large eggs x 50 g edible each (USDA), divided by 0.88 because shell is 12% of purchase weight.', 'Team 9')
fd['G' + str(FIRST + [f['id'] for f in FOODS].index('milk'))].comment = Comment('1 gallon = 16 cups x 244 g per cup (USDA SR28 portion weight).', 'Team 9')
fd['G' + str(FIRST + [f['id'] for f in FOODS].index('ice_cream'))].comment = Comment('1/2 gallon = 8 cups x 132 g per cup (USDA: 1/2 cup = 66 g).', 'Team 9')
fd['N' + str(HR)].comment = Comment('Palatability cap used as the upper bound u in the model (u = serving x max servings). These are team assumptions, not USDA data.', 'Team 9')
widths = {'A':4,'B':24,'C':13,'D':30,'E':9,'F':12,'G':11,'H':8,'I':10,'J':11,'K':8,'L':22,'M':10,'N':9,'Z':9,'AA':10,'AB':9,'AC':48,'AD':26,'AE':46}
for c, w in widths.items(): fd.column_dimensions[c].width = w
for n in NUTS: fd.column_dimensions[NUTCOL[n['key']]].width = 9
fd.freeze_panes = 'C6'

# ------------------------------------------------------------------ Matrix Model
mm = wb.create_sheet('Matrix Model')
put(mm, 'A1', 'Diet optimization as a linear program in matrix form', bold=True, size=14)
put(mm, 'A2', 'minimize  cᵀx    subject to    b_min ≤ A·x ≤ b_max,    0 ≤ x ≤ u        (x_j = amount of food j per day, in 100 g units)', bold=True)
put(mm, 'A3', 'A (11 nutrients × 36 foods) is pulled from Food Data.  A·x and cᵀx are matrix products (MMULT).  Change the blue targets, then re-solve (website optimizer or Excel Solver — see Notes & Sources).', italic=True)

# 1. targets
put(mm, 'A5', '1. TARGETS  (inputs — paste from the FeasaBULL calculator)', bold=True, size=11)
for ref, t in (('A6','Input'),('B6','Value / min'),('C6','Max'),('D6','Basis')): header(mm, ref, t)
T = [  # row, label, min, max, fmt, note
 (7,'Calorie target (kcal/day)',2600,None,'#,##0','Estimated Energy Requirement from the calculator'),
 (8,'Calorie tolerance (±)',0.05,None,'0%','Allowed band around the calorie target'),
 (9,'Protein (% of calories)',0.10,0.35,'0%','AMDR, adults'),
 (10,'Carbohydrate (% of calories)',0.45,0.65,'0%','AMDR, adults'),
 (11,'Fat (% of calories)',0.20,0.35,'0%','AMDR, adults'),
 (12,'Fiber (g)',38,None,'0','AI, men 19–50 (women: 25)'),
 (13,'Calcium (mg)',1000,None,'#,##0','RDA, adults 19–50'),
 (14,'Iron (mg)',8,None,'0','RDA, men 19+ (women 19–50: 18)'),
 (15,'Potassium (mg)',3400,None,'#,##0','AI, men 19+ (women: 2,600)'),
 (16,'Vitamin C (mg)',90,None,'0','RDA, men (women: 75)'),
 (17,'Sodium (mg)',None,2300,'#,##0','Chronic Disease Risk Reduction intake (upper)'),
 (18,'Total sugars (g)',None,None,'0','Optional maximum — leave blank for none'),
]
for r, lab, lo, hi, fmt, note in T:
    put(mm, f'A{r}', lab, border=True)
    put(mm, f'B{r}', lo, BLUE, fmt=fmt, border=True)
    put(mm, f'C{r}', hi, BLUE, fmt=fmt, border=True)
    put(mm, f'D{r}', note, border=True, italic=True)
put(mm, 'A19', 'Example person: male, 19–30, active, maintain weight. Replace these with the calculator output for each user.', italic=True)

# 2. constraint system
CR = 22  # header row
put(mm, f'A{CR-1}', '2. CONSTRAINT SYSTEM   b_min ≤ A·x ≤ b_max   (A = nutrient matrix: rows = nutrients, columns = foods, per 100 g edible)', bold=True, size=11)
for ref, t in (('A','i'),('B','Nutrient'),('C','Unit'),('D','b_min'),('E','b_max'),('F','A·x*  (optimized)'),('G','Status'),('H','A·x  (typical day)'),('I','Typical-day status')):
    header(mm, f'{ref}{CR}', t)
AC0 = 11  # column K
ACOL = [L(AC0 + j) for j in range(NF)]
AFIRST, ALAST = ACOL[0], ACOL[-1]
for j, f in enumerate(FOODS):
    c = put(mm, f'{ACOL[j]}{CR}', f"='Food Data'!B{FIRST + j}", GREEN, True, fill=SUB, align=CENTER, border=True)
mm.row_dimensions[CR].height = 66
R0 = CR + 1
RN = R0 + len(NUTS) - 1       # 23..33
BMIN = {'kcal':'=$B$7*(1-$B$8)','protein':'=$B$7*$B$9/4','carb':'=$B$7*$B$10/4','fat':'=$B$7*$B$11/9','fiber':'=$B$12','sugar':0,
        'calcium':'=$B$13','iron':'=$B$14','potassium':'=$B$15','sodium':0,'vitc':'=$B$16'}
BMAX = {'kcal':'=$B$7*(1+$B$8)','protein':'=$B$7*$C$9/4','carb':'=$B$7*$C$10/4','fat':'=$B$7*$C$11/9','fiber':1000000,'sugar':'=IF($C$18="",1000000,$C$18)',
        'calcium':1000000,'iron':1000000,'potassium':1000000,'sodium':'=$C$17','vitc':1000000}
ROW = {}
XR0 = 40; XRN = XR0 + NF - 1   # diet vector rows 40..75
for i, n in enumerate(NUTS):
    r = R0 + i; ROW[n['key']] = r
    put(mm, f'A{r}', i + 1, border=True)
    put(mm, f'B{r}', n['label'], bold=True, border=True)
    put(mm, f'C{r}', n['unit'], border=True)
    put(mm, f'D{r}', BMIN[n['key']], fmt='#,##0.0;-#,##0.0;"none"', border=True)
    put(mm, f'E{r}', BMAX[n['key']], fmt='[>=999999]"none";#,##0.0', border=True)
    for col, xc in (('G','F'),('I','H')):
        put(mm, f'{col}{r}', f'=IF({xc}{r}<D{r}-0.01,"BELOW MIN",IF({xc}{r}>E{r}+0.01,"ABOVE MAX",IF(OR(ABS({xc}{r}-D{r})<0.01,ABS({xc}{r}-E{r})<0.01),"binding","ok")))', border=True, align=CENTER)
    for j in range(NF):
        put(mm, f'{ACOL[j]}{r}', f"='Food Data'!{NUTCOL[n['key']]}{FIRST + j}", GREEN, fmt='0.00', border=True)
mm[f'F{R0}'] = ArrayFormula(f'F{R0}:F{RN}', f'=MMULT({AFIRST}{R0}:{ALAST}{RN},C{XR0}:C{XRN})')
mm[f'H{R0}'] = ArrayFormula(f'H{R0}:H{RN}', f'=MMULT({AFIRST}{R0}:{ALAST}{RN},I{XR0}:I{XRN})')
for r in range(R0, RN + 1):
    for col in 'FH':
        c = mm[f'{col}{r}']; c.font = font(bold=(col == 'F')); c.number_format = '#,##0.0'; c.border = BOX; c.fill = GREY
CROW, UROW = RN + 1, RN + 2   # 34, 35
put(mm, f'B{CROW}', 'c  — cost, $ per 100 g edible', bold=True, fill=SUB, border=True)
put(mm, f'B{UROW}', 'u  — cap, 100 g units / day', bold=True, fill=SUB, border=True)
for j in range(NF):
    fr = FIRST + j
    put(mm, f'{ACOL[j]}{CROW}', f"=IF('Food Data'!J{fr}=\"\",0,'Food Data'!J{fr})", GREEN, fmt='$0.000', border=True, fill=SUB)
    put(mm, f'{ACOL[j]}{UROW}', f"=IF('Food Data'!E{fr}=\"\",0,'Food Data'!K{fr}*'Food Data'!N{fr}/100)", GREEN, fmt='0.000', border=True, fill=SUB)
put(mm, f'C{CROW}', 'cᵀ (row vector)', italic=True)
put(mm, f'C{UROW}', 'u = 0 when no price', italic=True)

# 3. diet vectors
put(mm, f'A{XR0-3}', '3. DIET VECTORS x  (column vectors, 100 g units per day)', bold=True, size=11)
put(mm, f'A{XR0-2}', 'x* = optimal solution from the Simplex method for the example targets (pasted values; re-solve after changing targets).  Typical day = a hand-built, non-optimized day for comparison (edit freely).', italic=True)
for ref, t in (('A','j'),('B','Food'),('C','x*  optimized'),('D','u  (cap)'),('E','grams'),('F','servings'),('G','cost ($)'),('H','at cap?'),('I','x  typical day'),('J','grams'),('K','cost ($)')):
    header(mm, f'{ref}{XR0-1}', t)
TYP = {'eggs':1.0,'bread':0.58,'milk':2.44,'chicken_breast':1.13,'rice':0.92,'ground_beef':1.13,'american':0.28,'bananas':1.18,
       'iceberg':0.72,'tomatoes':0.90,'chips':0.28,'cookies':0.258}
for j, f in enumerate(FOODS):
    r = XR0 + j; fr = FIRST + j
    put(mm, f'A{r}', j + 1, border=True)
    put(mm, f'B{r}', f"='Food Data'!B{fr}", GREEN, border=True)
    put(mm, f'C{r}', round(xopt.get(f['id'], 0.0), 6), BLUE, fmt='0.0000', border=True, fill=GREY)
    put(mm, f'D{r}', f'=INDEX(${AFIRST}${UROW}:${ALAST}${UROW},1,A{r})', fmt='0.000', border=True)
    put(mm, f'E{r}', f'=C{r}*100', fmt='0', border=True)
    put(mm, f'F{r}', f"=E{r}/'Food Data'!K{fr}", fmt='0.00', border=True)
    put(mm, f'G{r}', f'=C{r}*INDEX(${AFIRST}${CROW}:${ALAST}${CROW},1,A{r})', fmt='$0.00', border=True)
    put(mm, f'H{r}', f'=IF(AND(D{r}>0,ABS(C{r}-D{r})<0.0001),"at cap","")', border=True, align=CENTER)
    put(mm, f'I{r}', TYP.get(f['id'], 0), BLUE, fmt='0.000', border=True)
    put(mm, f'J{r}', f'=I{r}*100', fmt='0', border=True)
    put(mm, f'K{r}', f'=I{r}*INDEX(${AFIRST}${CROW}:${ALAST}${CROW},1,A{r})', fmt='$0.00', border=True)
TR = XRN + 1  # 76
put(mm, f'B{TR}', 'Total cost per day   cᵀx', bold=True, fill=SUB, border=True)
put(mm, f'B{TR+1}', 'Total cost per week   7·cᵀx', bold=True, fill=SUB, border=True)
put(mm, f'B{TR+2}', 'Targets missed', bold=True, fill=SUB, border=True)
mm[f'G{TR}'] = ArrayFormula(f'G{TR}', f'=MMULT({AFIRST}{CROW}:{ALAST}{CROW},C{XR0}:C{XRN})')
mm[f'K{TR}'] = ArrayFormula(f'K{TR}', f'=MMULT({AFIRST}{CROW}:{ALAST}{CROW},I{XR0}:I{XRN})')
for col in 'GK':
    c = mm[f'{col}{TR}']; c.font = font(bold=True); c.number_format = '$0.00'; c.border = BOX; c.fill = SUB
    put(mm, f'{col}{TR+1}', f'={col}{TR}*7', bold=True, fmt='$0.00', border=True, fill=SUB)
put(mm, f'G{TR+2}', f'=COUNTIF(G{R0}:G{RN},"BELOW MIN")+COUNTIF(G{R0}:G{RN},"ABOVE MAX")', bold=True, border=True, fill=SUB)
put(mm, f'K{TR+2}', f'=COUNTIF(I{R0}:I{RN},"BELOW MIN")+COUNTIF(I{R0}:I{RN},"ABOVE MAX")', bold=True, border=True, fill=SUB)
put(mm, f'C{TR}', 'optimized →', italic=True); put(mm, f'I{TR}', 'typical →', italic=True)

# 4. proof via square linear system
ids = [f['id'] for f in FOODS]
free = ['rice', 'pasta', 'chicken_legs', 'oranges']
capped = ['flour', 'beans', 'milk']
bind = [('kcal', 'D', 'Calories ≥ min'), ('fat', 'D', 'Fat ≥ min'), ('fiber', 'D', 'Fiber ≥ min'), ('vitc', 'D', 'Vitamin C ≥ min')]
P0 = TR + 5   # title row 81
put(mm, f'A{P0}', '4. WHY x* IS OPTIMAL — at a Simplex corner the tight constraints form a square linear system  B·x_B = r', bold=True, size=11)
put(mm, f'A{P0+1}', 'In x*, 4 foods are free (0 < x < u), 3 sit at their caps (x = u), and exactly 4 nutrient constraints are tight. Moving the capped foods to the right-hand side leaves 4 equations in 4 unknowns: '
    'B·x_B = r, with r = b_B − N·u_N.  Since det(B) ≠ 0, x_B = B⁻¹·r is unique — and it reproduces the Simplex answer.  (Valid for the example targets; re-derive the tight set if targets change.)', italic=True)
H = P0 + 3  # header row 84
header(mm, f'B{H}', 'Tight constraint')
for k, fid in enumerate(free): header(mm, f'{L(3+k)}{H}', f"B: {FOODS[ids.index(fid)]['name']}")
for k, fid in enumerate(capped): header(mm, f'{L(8+k)}{H}', f"N: {FOODS[ids.index(fid)]['name']}")
for ref, t in (('L','b_B'),('M','r = b_B − N·u_N'),('O','x_B = B⁻¹·r'),('P','x* (Simplex)'),('Q','difference')): header(mm, f'{ref}{H}', t)
mm.row_dimensions[H].height = 54
B0, BN = H + 1, H + 4
for i, (key, _, lab) in enumerate(bind):
    r = B0 + i; ar = ROW[key]
    put(mm, f'B{r}', lab, bold=True, border=True)
    for k, fid in enumerate(free):
        put(mm, f'{L(3+k)}{r}', f'={ACOL[ids.index(fid)]}{ar}', fmt='0.000', border=True)
    for k, fid in enumerate(capped):
        put(mm, f'{L(8+k)}{r}', f'={ACOL[ids.index(fid)]}{ar}', fmt='0.000', border=True)
    put(mm, f'L{r}', f'=D{ar}', fmt='#,##0.000', border=True)
    put(mm, f'P{r}', f'=C{XR0 + ids.index(free[i])}', fmt='0.000000', border=True)
    put(mm, f'Q{r}', f'=O{r}-P{r}', fmt='0.0E+00', border=True)
# u_N column vector
header(mm, f'S{H}', 'Capped food'); header(mm, f'T{H}', 'u_N')
for k, fid in enumerate(capped):
    put(mm, f'S{B0+k}', FOODS[ids.index(fid)]['name'], border=True)
    put(mm, f'T{B0+k}', f'={ACOL[ids.index(fid)]}{UROW}', fmt='0.000', border=True)
mm[f'M{B0}'] = ArrayFormula(f'M{B0}:M{BN}', f'=L{B0}:L{BN}-MMULT(H{B0}:J{BN},T{B0}:T{B0+2})')
mm[f'O{B0}'] = ArrayFormula(f'O{B0}:O{BN}', f'=MMULT(C{B0+6}:F{BN+6},M{B0}:M{BN})')
for r in range(B0, BN + 1):
    for col, fmt in (('M', '#,##0.000'), ('O', '0.000000')):
        c = mm[f'{col}{r}']; c.number_format = fmt; c.border = BOX; c.font = font(bold=(col == 'O')); c.fill = GREY
# B inverse and determinant
IV0 = B0 + 6  # rows 91..94
put(mm, f'B{IV0-1}', 'B⁻¹  (MINVERSE)', bold=True)
put(mm, f'H{IV0-1}', 'det(B)  (MDETERM)', bold=True)
mm[f'C{IV0}'] = ArrayFormula(f'C{IV0}:F{IV0+3}', f'=MINVERSE(C{B0}:F{BN})')
for r in range(IV0, IV0 + 4):
    put(mm, f'B{r}', f'=B{B0 + r - IV0}', border=True)
    for col in 'CDEF':
        c = mm[f'{col}{r}']; c.number_format = '0.000000'; c.border = BOX; c.font = font(); c.fill = GREY
put(mm, f'H{IV0}', f'=MDETERM(C{B0}:F{BN})', bold=True, fmt='#,##0.0', border=True, fill=GREY)
put(mm, f'H{IV0+1}', f'=IF(ABS(H{IV0})>0.000001,"det ≠ 0 → B is invertible, unique solution","det = 0 → singular")', italic=True)
put(mm, f'L{IV0}', 'Check:', bold=True)
put(mm, f'M{IV0}', f'=IF(MAX(Q{B0}:Q{BN})-MIN(Q{B0}:Q{BN})<0.000001,"x_B = B⁻¹r matches the Simplex x*","MISMATCH — re-derive tight set")', bold=True)

# shadow prices
SP = IV0 + 6  # 97
put(mm, f'A{SP-1}', '5. SHADOW PRICES  yᵀ = c_Bᵀ·B⁻¹  — the extra $/day caused by raising each tight target by one unit', bold=True, size=11)
header(mm, f'B{SP}', 'Row vector')
for k, (key, _, lab) in enumerate(bind): header(mm, f'{L(3+k)}{SP}', lab)
put(mm, f'B{SP+1}', 'c_Bᵀ  (cost of free foods, $/100 g)', bold=True, border=True)
for k, fid in enumerate(free): put(mm, f'{L(3+k)}{SP+1}', f'={ACOL[ids.index(fid)]}{CROW}', fmt='$0.0000', border=True)
put(mm, f'B{SP+2}', 'yᵀ = c_Bᵀ·B⁻¹  ($ per unit)', bold=True, border=True)
mm[f'C{SP+2}'] = ArrayFormula(f'C{SP+2}:F{SP+2}', f'=MMULT(C{SP+1}:F{SP+1},C{IV0}:F{IV0+3})')
for col in 'CDEF':
    c = mm[f'{col}{SP+2}']; c.number_format = '$0.00000'; c.border = BOX; c.font = font(bold=True); c.fill = GREY
put(mm, f'B{SP+3}', 'Same, in friendlier units', bold=True, border=True)
for k, (mult, unit) in enumerate(((100, 'per +100 kcal'), (10, 'per +10 g fat'), (10, 'per +10 g fiber'), (10, 'per +10 mg vit C'))):
    put(mm, f'{L(3+k)}{SP+3}', f'=TEXT({L(3+k)}{SP+2}*{mult},"$0.000")&" {unit}"', border=True)
put(mm, f'B{SP+4}', 'Fat and fiber are the expensive constraints: they are what pushes the cheapest diet toward chicken legs and oranges instead of more rice.', italic=True)

# reduced costs / optimality certificate
RC = SP + 7  # 104
put(mm, f'A{RC-1}', '6. OPTIMALITY CERTIFICATE   d = c − yᵀ·A_B   (reduced cost of every food)', bold=True, size=11)
put(mm, f'B{RC-1+0}', None)
for j in range(NF): put(mm, f'{ACOL[j]}{RC}', f'={ACOL[j]}{CR}', GREEN, True, fill=SUB, align=CENTER, border=True)
mm.row_dimensions[RC].height = 66
put(mm, f'B{RC}', 'A_B  = tight rows of A', bold=True)
for i, (key, _, lab) in enumerate(bind):
    r = RC + 1 + i
    put(mm, f'B{r}', lab, border=True)
    for j in range(NF): put(mm, f'{ACOL[j]}{r}', f'={ACOL[j]}{ROW[key]}', fmt='0.00', border=True)
DR = RC + 5
put(mm, f'B{DR}', 'd = c − yᵀ·A_B   ($ per 100 g)', bold=True, border=True, fill=SUB)
mm[f'{AFIRST}{DR}'] = ArrayFormula(f'{AFIRST}{DR}:{ALAST}{DR}', f'={AFIRST}{CROW}:{ALAST}{CROW}-MMULT(C{SP+2}:F{SP+2},{AFIRST}{RC+1}:{ALAST}{RC+4})')
for j in range(NF):
    c = mm[f'{ACOL[j]}{DR}']; c.number_format = '0.0000'; c.border = BOX; c.font = font(bold=True); c.fill = GREY
put(mm, f'B{DR+1}', 'Condition (x = 0 → d ≥ 0;  at cap → d ≤ 0;  free → d = 0)', bold=True, border=True)
for j in range(NF):
    col = ACOL[j]
    xr = f'INDEX($C${XR0}:$C${XRN},{j+1})'; ur = f'{col}${UROW}'
    put(mm, f'{col}{DR+1}', f'=IF({ur}=0,"no price",IF({xr}<0.000001,IF({col}{DR}>=-0.000001,"OK","FAIL"),IF(ABS({xr}-{ur})<0.000001,IF({col}{DR}<=0.000001,"OK","FAIL"),IF(ABS({col}{DR})<0.000001,"OK","FAIL"))))', border=True, align=CENTER)
put(mm, f'B{DR+2}', 'Is x* provably optimal?', bold=True, border=True, fill=SUB)
put(mm, f'C{DR+2}', f'=IF(COUNTIF({AFIRST}{DR+1}:{ALAST}{DR+1},"FAIL")=0,"YES — no food can be added or swapped to lower the cost","NO — re-solve")', bold=True)

mm.column_dimensions['A'].width = 28
mm.column_dimensions['B'].width = 30
for c in 'CDEFGHIJ': mm.column_dimensions[c].width = 14
for c in ACOL: mm.column_dimensions[c].width = max(mm.column_dimensions[c].width or 0, 11)
mm.freeze_panes = 'C23'

# ------------------------------------------------------------------ Notes & Sources
ns = wb.create_sheet('Notes & Sources')
ns.column_dimensions['A'].width = 140
lines = [
 ('How this workbook maps to the project', True),
 ('Food Data = the raw data. Matrix Model = the linear algebra: nutrient matrix A, cost vector c, target vectors b_min / b_max, diet vector x, and the products A·x and cᵀx (MMULT).', False),
 ('Section 4 proves the answer with a square linear system (MINVERSE, MDETERM); sections 5–6 give shadow prices and the optimality certificate.', False),
 ('The website optimizer (diet-optimizer.js) uses exactly this A, c, b and solves it with a two-phase Simplex written as Gauss-Jordan row operations.', False),
 ('', False),
 ('Re-solving inside Excel (Solver add-in, Data → Solver)', True),
 ("Set Objective: 'Matrix Model'!G76   To: Min   By Changing Variable Cells: 'Matrix Model'!C40:C75", False),
 ("Constraints: F23:F33 >= D23:D33;  F23:F33 <= E23:E33;  C40:C75 <= D40:D75.   Tick 'Make Unconstrained Variables Non-Negative'.   Solving Method: Simplex LP.", False),
 ('Google Sheets: install the free "Solver" add-on, or paste x from the website. Sections 4–6 only apply to the example targets until you re-derive the tight constraints.', False),
 ('', False),
 ('Assumptions to state in the presentation', True),
 ('1. Prices are national averages (BLS U.S. city average), not Tampa store prices. Regional prices differ; a store-specific price list can replace column E of Food Data.', False),
 ('2. Nutrients are per 100 g of the food as bought and edible (raw meat, dry rice/pasta/beans). Cooking mostly changes water weight; small cooking losses (e.g., vitamin C, fat dripping from meat) are ignored.', False),
 ('3. Prices are converted to $ per 100 g EDIBLE portion with USDA refuse percentages (bone in chicken, peel on bananas/oranges, shells on eggs). Without this, bone-in chicken looks ~30% cheaper than it is.', False),
 ('4. BLS item → USDA food matches are judgment calls where BLS is generic: "Ground beef, 100% beef" = 80/20, "lean and extra lean" = 90/10, "Beans, dried, any type" = pinto.', False),
 ('5. "Max servings / day" caps are team assumptions to keep the diet realistic. Without caps, a pure min-cost LP eats implausible amounts of the cheapest foods (Stigler, 1945; Dantzig, 1990).', False),
 ('6. BLS discontinued average prices for peanut butter (Dec 2017), apples (Oct 2017), broccoli (Mar 2020), carrots (Mar 2014) and butter (Apr 2012), and never priced oats or canned tuna. Those rows need a dated local price before they enter the model.', False),
 ('7. With only BLS-priced foods, a vegan plan is infeasible: no priced vegan food supplies enough fat or calcium (no oil, peanut butter, tofu, or fortified soy milk). Add priced foods to fix this.', False),
 ('8. USDA SR28 is from 2015; FoodData Central\'s "SR Legacy" (2018) carries the same NDB numbers and is where these foods can be looked up online.', False),
 ('', False),
 ('References (APA 7)', True),
 ('U.S. Bureau of Labor Statistics. (2026). Average retail food and energy prices, U.S. and Midwest region [Data table]. Retrieved October 26, 2026, from https://www.bls.gov/regions/mid-atlantic/data/averageretailfoodandenergyprices_usandmidwest_table.htm', False),
 ('U.S. Department of Agriculture, Agricultural Research Service, Nutrient Data Laboratory. (2015). USDA National Nutrient Database for Standard Reference, Release 28 [Data set]. http://www.ars.usda.gov/nea/bhnrc/ndl', False),
 ('U.S. Department of Agriculture, Agricultural Research Service. (2019). FoodData Central: SR Legacy [Data set]. https://fdc.nal.usda.gov/', False),
 ('U.S. Bureau of Labor Statistics. (2019). Average price: Peanut butter, creamy, all sizes (cost per pound/453.6 grams) in U.S. city average [APU0000716141; discontinued]. Federal Reserve Bank of St. Louis. https://fred.stlouisfed.org/series/APU0000716141', False),
 ('Institute of Medicine. (2005). Dietary reference intakes for energy, carbohydrate, fiber, fat, fatty acids, cholesterol, protein, and amino acids. The National Academies Press. https://doi.org/10.17226/10490', False),
 ('National Academies of Sciences, Engineering, and Medicine. (2019). Dietary reference intakes for sodium and potassium. The National Academies Press. https://doi.org/10.17226/25353', False),
 ('Dantzig, G. B. (1990). The diet problem. Interfaces, 20(4), 43–47. https://doi.org/10.1287/inte.20.4.43', False),
 ('Stigler, G. J. (1945). The cost of subsistence. Journal of Farm Economics, 27(2), 303–314. https://doi.org/10.2307/1231810', False),
 ('', False),
 ('Color legend', True),
 ('Blue text = input data you can edit.  Black = formula.  Green = link to another sheet.  Yellow fill = price still needed.  Grey fill = matrix-product results.', False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ns.cell(row=i, column=1, value=t); c.font = font(bold=b, size=11 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')

wb.save('Team9_Diet_Optimization_Data.xlsx')
print('saved; rows', FIRST, LAST, 'matrix rows', R0, RN, 'x rows', XR0, XRN, 'total row', TR, 'proof', P0, B0, IV0, SP, RC, DR)
