"""Exit 0 if two workbooks have the same sheets, formulas and values (ignores timestamps and styles), else 1.
Used by CI so the spreadsheet is only re-committed when its content actually changed.
    python scripts/same_xlsx.py old.xlsx new.xlsx
"""
import sys
from openpyxl import load_workbook

def norm(v):
    if isinstance(v, float): return round(v, 9)
    if hasattr(v, 'text'): return f'{{{v.ref}}}{v.text}'      # ArrayFormula: compare its range and text, not the object
    return str(v)

def content(path):
    out = []
    for data_only in (False, True):
        for ws in load_workbook(path, data_only=data_only):
            out.append((ws.title, [(c.coordinate, norm(c.value)) for row in ws.iter_rows() for c in row if c.value is not None]))
    return out

try:
    same = content(sys.argv[1]) == content(sys.argv[2])
except Exception as e:                       # missing/unreadable old file counts as changed
    print(f'cannot compare ({e}); treating as changed'); same = False
print('workbook content unchanged' if same else 'workbook content changed')
sys.exit(0 if same else 1)
