"""Recalculate the workbook with LibreOffice (headless) so every formula has a stored value, then check for errors.

    python scripts/recalc_xlsx.py [spreadsheet/FeasaBULL_Diet_Optimizer.xlsx]

openpyxl writes formulas without results. Excel recalculates on open, but GitHub previews, Google Sheets imports
and phones show blanks until a recalc. Exits 1 (and keeps the original file) if any formula returns an error.
"""
import os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
ERRORS = ('#VALUE!', '#DIV/0!', '#REF!', '#NAME?', '#N/A', '#NUM!', '#NULL!')

def main(path):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        sys.exit('LibreOffice (soffice) not found. Install it, e.g. apt-get install -y libreoffice-calc')
    with tempfile.TemporaryDirectory() as tmp:
        profile = 'file://' + os.path.join(tmp, 'profile')            # private profile: no clash with a running LibreOffice
        subprocess.run([soffice, f'-env:UserInstallation={profile}', '--headless', '--calc', '--convert-to', 'xlsx',
                        '--outdir', tmp, path], check=True, capture_output=True, timeout=300)
        out = os.path.join(tmp, os.path.basename(path))
        n_before = sum(1 for ws in load_workbook(path) for row in ws.iter_rows() for c in row if c.data_type == 'f')
        vals = load_workbook(out, data_only=True)
        n_after = sum(1 for ws in load_workbook(out) for row in ws.iter_rows() for c in row if c.data_type == 'f')
        bad = [f'{ws.title}!{c.coordinate} = {c.value}' for ws in vals for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value in ERRORS]
        if bad:
            print(f'{len(bad)} formula errors (original kept):\n  ' + '\n  '.join(bad[:20])); sys.exit(1)
        if n_after < n_before:
            print(f'LibreOffice dropped formulas ({n_before} -> {n_after}); original kept.'); sys.exit(1)
        shutil.copyfile(out, path)
        print(f'recalculated {os.path.relpath(path, ROOT)}: {n_after} formulas, 0 errors')

if __name__ == '__main__':
    main(os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'spreadsheet', 'FeasaBULL_Diet_Optimizer.xlsx')))
